"""Makes the real DL / CDL list, with UPMC circle and division names, in private/ (never published):
  private/data.json            for checks and tests
  private/CDL_BS_import.xlsx   the starting Google Sheet (needs openpyxl): upload to Drive and open with Google Sheets
Run: python3 scripts/make_data.py <CDL/Cleaning folder>"""
import sys, os, json
from collections import Counter
sys.path.insert(0, os.path.dirname(__file__))
import duty_classify as K

K.SRC = sys.argv[1]
PRIVATE = os.path.join(os.path.dirname(__file__), '..', 'private')

# cleaned-data circle -> UPMC circle
CIRCLE = {'P&G': 'Parks & Gardens', 'SC (North)': 'Soil Con (N)', 'WL (HQ)': 'WL HQ', 'WL North': 'Wild Life North'}
# cleaned-data division -> UPMC division ("... Circle Office" becomes "<circle> Circle HQ" so names stay unique)
DIVISION = {'Utilisation': 'Utilization', 'BTR (EAST)': 'B.T.R (East)', 'BTR (West)': 'B.T.R (West)', 'Bankura (North)': 'Bankura (N)', 'Bankura (South)': 'Bankura (S)',
 'SFTI, Hijli': 'SFTI Hizli', 'WBFS, Dow Hill': 'WB Forest School', 'Kurseong': 'Kurseong Dn.', 'Monitoring (North)': 'Monitoring North', 'Monitoring (South)': 'Monitoring South',
 'Jalpaiguri SF': 'Jalpaiguri S.F.', 'Siliguri SF': 'Siliguri S.F', 'BKP': 'Baikunthapur', 'Cooch Behar': 'Coochbehar', 'P&G (North)': 'Parks & Gardens North',
 'Silviculture (Hills)': 'Silviculture Hill', 'Silviculture (North)': 'Silviculture North', 'Silviculture (South)': 'Silviculture South',
 '24-Parganas (N)': '24 Parganas (N)', '24-Parganas (S)': '24 Parganas (S)', 'SBR HQ': 'SBR Circle HQ', 'STR': 'S.T.R', 'Extension Forestry': 'Extn. Forestry',
 'Kangsabati North': 'Kangsabati (N)', 'Kangsabati South': 'Kangsabati (S)', 'Wildlife (HQ)': 'WL HQ Division', 'WP & GIS (North)': 'WP North',
 'WP & GIS (South-I)': 'WP South I', 'WP & GIS (South-II)': 'WP South II'}

DUTY = {  # sub-category from the classifier -> one of the 9 duty categories
 'Cleaning / sweeping': 'Cleaning/Sweeping/Maintenance', 'Maintenance & utilities': 'Cleaning/Sweeping/Maintenance', 'Rest house / bungalow': 'Cleaning/Sweeping/Maintenance',
 'Nursery': 'Nursery Operations', 'Plantation work': 'Nursery Operations', 'Timber depot': 'Timber Depot', 'Driver': 'Driver', 'Garden / park': 'Mali/Gardener',
 'Patrolling & forest protection': 'Forest Protection and Patrolling', 'Plantation watch & ward': 'Forest Protection and Patrolling',
 'Check post / gate': 'Forest Protection and Patrolling', 'Fire / watch tower': 'Forest Protection and Patrolling',
 'Elephant / wildlife conflict': 'Human-Animal Conflict Management', 'Wildlife rescue & release': 'Human-Animal Conflict Management',
 'Night guard / security': 'Night Guard/Security'}
PLACE = {'Beat': 'Beat', 'Range HQ': 'Range office', 'Range (beat not stated)': 'Range office', 'Division HQ': 'Division office', 'Timber depot': 'Timber depot', 'Rest house': 'Rest house'}

def date(r):
    d, p = r['Engagement_Date'], r['Date_Precision']
    if not d: return ''
    return str(d.year) if p == 'Year' else d.strftime('%d-%m-%Y')

out = []
for r in K.load():
    circle = CIRCLE.get(r['Circle'], r['Circle'])
    div = r['Division']
    div = circle + ' Circle HQ' if div.endswith('Circle Office') else DIVISION.get(div, div)
    duty = '' if r['Cat'] in ('Unclear', 'Not reported') else DUTY.get(r['Sub'], 'Others')
    other = []
    for s in r['Also']:
        c = DUTY.get(s, 'Others')
        if c != duty and c not in other: other.append(c)
    place = PLACE.get(r['Posting_Location_Type'], '')
    pname = (r['Other_Place'] or '') if place not in ('', 'Beat') else ''
    mode = {'Direct': 'Direct', 'Through Agency': 'Through agency'}.get(r['Engagement_Mode'], '')
    out.append({'id': r['Record_ID'], 'circle': circle, 'division': div, 'name': (r['Name'] or '').strip(), 'mode': mode, 'agency': '',
                'date': date(r), 'district': '' if r['District'] in (None, 'Not stated') else r['District'],
                'range': 'Division HQ' if place == 'Division office' else '' if r['Range'] in (None, 'Not stated') else r['Range'],
                'beat': '' if place == 'Division office' else r['Beat'] or '',
                'place': place, 'placeName': pname, 'duty': duty, 'other': other,
                'pay': r['Monthly_Pay_Jul2026'], 'details': '' if duty != 'Others' else (r['Resolved'] or '')[:80], 'flag': ''})

# Possible double entry: same name and same engagement date twice in one division
key = lambda o: (o['division'], ' '.join(o['name'].lower().split()), o['date'])
twice = Counter(key(o) for o in out if o['name'] and o['date'])
for o in out:
    if twice.get(key(o), 0) > 1:
        o['flag'] = 'Possible double entry: same name and engagement date appear twice in this division.'

os.makedirs(PRIVATE, exist_ok=True)
with open(os.path.join(PRIVATE, 'data.json'), 'w') as f:
    json.dump(out, f, ensure_ascii=False)

# ---- the import workbook: the same tabs and headings as apps-script/Code.gs (COLS, LOG_HEADERS, CERT_HEADERS)
COLS = [('id', 'ID'), ('circle', 'Circle'), ('division', 'Division'), ('name', 'Name'), ('mode', 'Mode of engagement'), ('agency', 'Agency name'),
        ('date', 'Engagement date'), ('district', 'District'), ('range', 'Range (or Division HQ)'), ('beat', 'Beat'), ('place', 'Place of work'),
        ('placeName', 'Name of the place'), ('duty', 'Main duty'), ('other', 'Other duties'), ('details', 'Duty details'), ('pay', 'Monthly pay (Rs)'),
        ('remarks', 'Remarks'), ('ver', 'Verification'), ('state', 'Status'), ('flag', 'Problem'), ('removal', 'Removal request'),
        ('updOn', 'Last updated on'), ('updBy', 'Last updated by'), ('orig', 'As in the return (do not edit)')]
TRACK = ['name', 'mode', 'agency', 'date', 'district', 'range', 'beat', 'place', 'placeName', 'duty', 'other', 'details', 'pay', 'remarks']
LOG_HEADERS = ['When', 'Username', 'By', 'Circle', 'Division', 'Cadre', 'ID', 'Person', 'Action', 'Field', 'Old value', 'New value']
CERT_HEADERS = ['Division', 'Cadre', 'Certified by', 'Designation', 'On', 'Entered by']
# UPMC / FV circles and divisions, so divisions with no DL / CDL can still enter Ban Sahayaks
UPMC = [('APM', ['Circle HQ', 'Admin', 'Utilization', 'Publicity']), ('WP & GIS', ['Circle HQ', 'WP North', 'WP South I', 'WP South II']),
        ('Central', ['Circle HQ', 'Bankura (N)', 'Bankura (S)', 'Panchet']), ('Western', ['Circle HQ', 'Medinipur', 'Purba Medinipur', 'Kharagpur', 'Jhargram', 'Rupnarayan']),
        ('South East', ['Circle HQ', 'Burdwan', 'Birbhum', 'Durgapur']), ('South West', ['Circle HQ', 'Purulia', 'Kangsabati (N)', 'Kangsabati (S)', 'Extn. Forestry']),
        ('Parks & Gardens', ['Circle HQ', 'Howrah', 'Parks & Gardens North', 'URF']), ('North West', ['Circle HQ', 'Malda', 'Raiganj', 'Siliguri S.F', 'Jalpaiguri S.F.']),
        ('Hill', ['Circle HQ', 'Darjeeling', 'Kalimpong', 'Kurseong Dn.']), ('Northern', ['Circle HQ', 'Jalpaiguri', 'Baikunthapur', 'NTFP', 'Coochbehar']),
        ('Wild Life North', ['Circle HQ', 'Darjeeling WL', 'Gorumara WL', 'Jaldapara WL']), ('BTR', ['Circle HQ', 'B.T.R (East)', 'B.T.R (West)']),
        ('WL HQ', ['WL HQ Division']), ('STR', ['S.T.R']), ('SBR', ['Circle HQ', '24 Parganas (S)', '24 Parganas (N)', 'Nadia-Murshidabad']),
        ('Research', ['Circle HQ', 'Silviculture Hill', 'Silviculture North', 'Silviculture South']), ('Development', ['Circle HQ', 'WB Forest School', 'SFTI Hizli']),
        ('Monitoring', ['Circle HQ', 'Monitoring North', 'Monitoring South']), ('Soil Con (N)', ['Circle HQ', 'Kurseong SC', 'Jalpaiguri SC'])]
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:
    Workbook = None
    print('openpyxl not installed: data.json written, workbook skipped')
if Workbook:
    wb = Workbook()
    def tab(name, headers, rows, first=False):
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        ws.append(headers)
        for c in ws[1]: c.font = Font(bold=True)
        ws.freeze_panes = 'A2'
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):   # plain text everywhere, so "2009" and "02-09-2009" stay as written
            for c in row: c.number_format = '@'
        return ws
    def cell(o, k):
        if k == 'other': return '; '.join(o['other'])
        if k == 'state': return 'Working'
        if k == 'orig': return json.dumps({f: o.get(f, [] if f == 'other' else '') if o.get(f) is not None else ('' if f != 'pay' else None) for f in TRACK}, ensure_ascii=False)
        v = o.get(k, '')
        return '' if v is None else (str(v) if k != 'pay' else v)
    tab('DL-CDL', [h for _, h in COLS], [[cell(o, k) for k, _ in COLS] for o in out], first=True)
    tab('Ban Sahayak', [h for _, h in COLS], [])
    tab('Change Log', LOG_HEADERS, [])
    tab('Certificates', CERT_HEADERS, [])
    lists, seen = [], set()
    for c, ds in UPMC:
        for d in ds:
            d = c + ' Circle HQ' if d == 'Circle HQ' else d
            if d not in seen: seen.add(d); lists.append([c, d])
    for o in out:
        if o['division'] not in seen: seen.add(o['division']); lists.append([o['circle'], o['division']])
    tab('Lists', ['Circle', 'Division'], lists)
    wb.save(os.path.join(PRIVATE, 'CDL_BS_import.xlsx'))
    print('workbook:', len(out), 'DL / CDL rows,', len(lists), 'divisions in Lists')

print(len(out), 'people,', len({o['division'] for o in out}), 'divisions,', sum(1 for o in out if o['flag']), 'flagged')
print(Counter(o['circle'] for o in out))
