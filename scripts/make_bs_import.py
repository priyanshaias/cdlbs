"""Makes private/BS_import.xlsx: the Ban Sahayaks from "FINAL 1500 LIST OF BANASAHAYAKS.xlsx", in the columns of the
Ban Sahayak tab of the CDL+BS sheet (append it there; see the README). Real names: the file stays in private/.
Run: python3 scripts/make_bs_import.py "<path to FINAL 1500 LIST OF BANASAHAYAKS.xlsx>"   (needs openpyxl)"""
import sys, os, re, json, secrets
from collections import Counter
from openpyxl import load_workbook, Workbook

SRC = sys.argv[1]
OUT = os.path.join(os.path.dirname(__file__), '..', 'private', 'BS_import.xlsx')

# the list's office -> (circle, division) as in the CDL+BS Lists tab
DIV = {
 'D.C.F - Administration': ('APM', 'Admin'), 'D.C.F - Publicity': ('APM', 'Publicity'), 'D.F.O - Forest Utilization': ('APM', 'Utilization'),
 'Working Plan North - Darjeeling': ('WP & GIS', 'WP North'), 'Working Plan South - I': ('WP & GIS', 'WP South I'), 'Working Plan South - II': ('WP & GIS', 'WP South II'),
 'D.F.O - Jhargram': ('Western', 'Jhargram'), 'D.F.O - Medinipur': ('Western', 'Medinipur'), 'D.F.O - Kharagpur': ('Western', 'Kharagpur'),
 'D.F.O - Purba Medinipur': ('Western', 'Purba Medinipur'), 'D.F.O - Rupnarayan': ('Western', 'Rupnarayan'),
 'D.F.O - Purulia': ('South West', 'Purulia'), 'D.F.O - Kangsabati North': ('South West', 'Kangsabati (N)'), 'D.F.O - Kangsabati South': ('South West', 'Kangsabati (S)'),
 'D.F.O - Extension Forestry': ('South West', 'Extn. Forestry'),
 'D.F.O - Bankura North': ('Central', 'Bankura (N)'), 'D.F.O - Bankura South': ('Central', 'Bankura (S)'), 'D.F.O - Panchet': ('Central', 'Panchet'),
 'D.F.O - Durgapur': ('South East', 'Durgapur'), 'D.F.O.Birbhum': ('South East', 'Birbhum'), 'D.F.O. Burdwan': ('South East', 'Burdwan'),
 'CF-P & G Office': ('Parks & Gardens', 'Parks & Gardens Circle HQ'), 'D.F.O - Howrah': ('Parks & Gardens', 'Howrah'), 'D.C.F - U.R.F': ('Parks & Gardens', 'URF'),
 'D.F.O - P & G North': ('Parks & Gardens', 'Parks & Gardens North'),
 'CF-North West Office': ('North West', 'North West Circle HQ'), 'D.F.O - Siliguri S.F': ('North West', 'Siliguri S.F'), 'D.F.O - Raiganj': ('North West', 'Raiganj'),
 'D.F.O - Malda': ('North West', 'Malda'), 'D.F.O - Jalpaiguri S.F': ('North West', 'Jalpaiguri S.F.'),
 'CCF-Hill Office': ('Hill', 'Hill Circle HQ'), 'D.F.O - Kalimpong': ('Hill', 'Kalimpong'), 'D.F.O - Darjeeling': ('Hill', 'Darjeeling'), 'D.F.O - Kurseong': ('Hill', 'Kurseong Dn.'),
 'CCF-NC Office': ('Northern', 'Northern Circle HQ'), 'D.F.O - N.T.F.P': ('Northern', 'NTFP'), 'D.F.O - Baikunthapur': ('Northern', 'Baikunthapur'),
 'D.F.O - Jalpaiguri': ('Northern', 'Jalpaiguri'), 'D.F.O - Cooch Behar': ('Northern', 'Coochbehar'),
 'D.F.O - Wildlife Headquarter': ('WL HQ', 'WL HQ Division'),
 'CCF-WL N office': ('Wild Life North', 'Wild Life North Circle HQ'), 'D.F.O - Darjeeling Wildlife': ('Wild Life North', 'Darjeeling WL'),
 'D.F.O - Gorumara Wildlife': ('Wild Life North', 'Gorumara WL'), 'D.F.O - Jaldapara Wildlife': ('Wild Life North', 'Jaldapara WL'),
 'CCF & FD,BTR Office': ('BTR', 'BTR Circle HQ'), 'D.F.D - B.T.R West': ('BTR', 'B.T.R (West)'), 'D.F.D - B.T.R East': ('BTR', 'B.T.R (East)'),
 'DIR-SBR-Office': ('SBR', 'SBR Circle HQ'), 'D.F.O - 24 Pgs South': ('SBR', '24 Parganas (S)'), 'D.F.O - 24 Pgs North': ('SBR', '24 Parganas (N)'),
 'D.F.O - Nadia & Murshidabad': ('SBR', 'Nadia-Murshidabad'), 'D.F.D - S.T.R': ('STR', 'S.T.R'),
 'CCF-Monitoring Office': ('Monitoring', 'Monitoring Circle HQ'), 'D.F.O - Monitoring South': ('Monitoring', 'Monitoring South'),
 'D.F.O - Monitoring North': ('Monitoring', 'Monitoring North'), 'CF RESEARCH': ('Research', 'Research Circle HQ'),
 'D.F.O - Kurseong Soil Conv.': ('Soil Con (N)', 'Kurseong SC'), 'D.F.O - Jalpaiguri Soil Conv.': ('Soil Con (N)', 'Jalpaiguri SC'),
 'D.F.O - Silvi Hill': ('Research', 'Silviculture Hill'), 'D.F.O - Silvi North': ('Research', 'Silviculture North'), 'D.F.O - Silvi South': ('Research', 'Silviculture South'),
 'Director - W.B.F.S.': ('Development', 'WB Forest School'), 'Director - S.F.T.I': ('Development', 'SFTI Hizli'),
 'Zoological Garden Alipore': ('Zoo', 'Alipore Zoological Garden'),   # not in the Lists tab yet: add it there
}
# "CCF,Office" (a circle's own office) opens several circles' blocks: the next division listed says which circle
SKIP_NAMES = {'NIL', ''}

COLS = ['ID', 'Circle', 'Division', 'Name', 'Mode of engagement', 'Agency name', 'Engagement date', 'District', 'Range (or Division HQ)', 'Beat',
        'Place of work', 'Name of the place', 'Main duty', 'Other duties', 'Duty details', 'Monthly pay (Rs)', 'Remarks', 'Verification', 'Status',
        'Problem', 'Removal request', 'Last updated on', 'Last updated by', 'As in the return (do not edit)']

ws = load_workbook(SRC, data_only=True).worksheets[0]
rows = list(ws.iter_rows(min_row=3, values_only=True))
# circle of each "CCF,Office" block, from the next division in the list
office_circle, pending = {}, []
for i, r in enumerate(rows):
    d = str(r[0]).strip() if r[0] else ''
    if d == 'CCF,Office': pending.append(i)
    elif d in DIV and pending:
        for j in pending: office_circle[j] = DIV[d][0]
        pending = []

clean = lambda s: re.sub(r'\s+', ' ', str(s or '')).strip()
out, skipped, cur, cur_i = [], [], None, None
for i, r in enumerate(rows):
    d, _, name, gen, cas, de, ren = (list(r) + [None] * 7)[:7]
    if d and str(d).strip(): cur, cur_i = str(d).strip(), i
    name = clean(name).upper()
    if name in SKIP_NAMES: continue
    if cur.startswith('Deputy Conservator') or 'Conservator of Forests' in name: skipped.append((cur, name)); continue   # signature line
    circle, div = (office_circle[cur_i] + '', office_circle[cur_i] + ' Circle HQ') if cur == 'CCF,Office' else DIV[cur]
    m = re.match(r'^(\d{1,2})[./-](\d{1,2})[./-](\d{4})$', clean(de))
    date = '%02d-%02d-%s' % (int(m[1]), int(m[2]), m[3]) if m else ''
    g = clean(gen).replace(',', '').upper(); g = '' if g == 'NIL' else g
    c = clean(cas).upper(); c = '' if c == 'NIL' else c
    remarks = ', '.join(x for x in [('Gender ' + g) if g else '', ('caste ' + c) if c else ''] if x)
    out.append({'circle': circle, 'division': div, 'name': name, 'date': date, 'remarks': remarks[:1].upper() + remarks[1:] if remarks else ''})

# possible double entries: same name and date in one division
key = lambda o: (o['division'], o['name'], o['date'])
twice = Counter(key(o) for o in out)
ids = set()
def new_id():
    while True:
        x = 'BS-' + secrets.token_hex(4).upper()
        if x not in ids: ids.add(x); return x

# no heading row: the file is appended under the tab's own headings, in the same column order (COLS)
wb = Workbook(); sh = wb.active; sh.title = 'Ban Sahayak'
for o in out:
    flag = 'Possible double entry: same name and engagement date appear twice in this division.' if twice[key(o)] > 1 else ''
    row = {'ID': new_id(), 'Circle': o['circle'], 'Division': o['division'], 'Name': o['name'], 'Mode of engagement': 'Direct', 'Engagement date': o['date'],
           'Monthly pay (Rs)': '12000', 'Remarks': o['remarks'], 'Status': 'Working', 'Problem': flag,
           'Last updated by': 'Import: FINAL 1500 LIST OF BANASAHAYAKS'}
    sh.append([row.get(h, '') for h in COLS])
for r in sh.iter_rows(min_row=1):   # plain text, so dates stay as written
    for c in r: c.number_format = '@'
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)

by = Counter(o['division'] for o in out)
print(len(out), 'Ban Sahayaks in', len(by), 'divisions ->', os.path.normpath(OUT))
print('skipped (not people):', skipped)
print('no engagement date:', sum(1 for o in out if not o['date']), '· possible double entries:', sum(1 for o in out if twice[key(o)] > 1))
print(json.dumps(dict(sorted(by.items())), ensure_ascii=False))
